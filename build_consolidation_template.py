#!/usr/bin/env python3
"""Builds assets/downloads/multi-entity-consolidation-template.xlsx (lead magnet, 2026-10-24).
Run with the base Python that has openpyxl:  C:/Users/Admin/AppData/Local/Python/bin/python.exe build_consolidation_template.py
Sheets: Start Here, Group Lines, Company A/B/C (trial balances), Eliminations, Consolidated TB, P&L, Balance Sheet.
Sign convention: Net = Debit - Credit, so income and liabilities are negative in the TB and flipped on the statements.
Example figures for Company A/B are illustrative only (A charges B a 10,000 management fee)."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "downloads",
                   "multi-entity-consolidation-template.xlsx")
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F2937")
INPUT = PatternFill("solid", fgColor="FEF3C7")
TITLE = Font(bold=True, size=14)
NUM = '#,##0;(#,##0);"-"'
ROWS = 60  # input rows per trial balance

# group line -> (statement, section)
GROUP_LINES = [
    ("Revenue", "P&L", "Income"),
    ("Other income", "P&L", "Income"),
    ("Intercompany income", "P&L", "Income"),
    ("Cost of sales", "P&L", "Cost of sales"),
    ("Operating expenses", "P&L", "Overheads"),
    ("Depreciation", "P&L", "Overheads"),
    ("Intercompany expense", "P&L", "Overheads"),
    ("Cash at bank", "BS", "Current assets"),
    ("Trade receivables", "BS", "Current assets"),
    ("Intercompany receivable", "BS", "Current assets"),
    ("Fixed assets", "BS", "Non-current assets"),
    ("Trade payables", "BS", "Current liabilities"),
    ("Intercompany payable", "BS", "Current liabilities"),
    ("Tax payable", "BS", "Current liabilities"),
    ("Loans", "BS", "Non-current liabilities"),
    ("Share capital", "BS", "Equity"),
    ("Retained earnings", "BS", "Equity"),
]

COMPANY_A = [
    ("4000", "Sales", 0, 250000, "Revenue"),
    ("4900", "Management fees - Company B", 0, 10000, "Intercompany income"),
    ("5000", "Cost of goods sold", 120000, 0, "Cost of sales"),
    ("6000", "Wages and overheads", 60000, 0, "Operating expenses"),
    ("6500", "Depreciation", 5000, 0, "Depreciation"),
    ("1000", "Bank", 80000, 0, "Cash at bank"),
    ("1100", "Debtors", 40000, 0, "Trade receivables"),
    ("1150", "Due from Company B", 10000, 0, "Intercompany receivable"),
    ("1500", "Plant and equipment", 50000, 0, "Fixed assets"),
    ("2000", "Creditors", 0, 30000, "Trade payables"),
    ("2200", "Corporation tax", 0, 15000, "Tax payable"),
    ("3000", "Share capital", 0, 10000, "Share capital"),
    ("3100", "Retained earnings", 0, 50000, "Retained earnings"),
]
COMPANY_B = [
    ("200", "Revenue", 0, 120000, "Revenue"),
    ("310", "Purchases", 50000, 0, "Cost of sales"),
    ("400", "General expenses", 35000, 0, "Operating expenses"),
    ("410", "Management fee - Company A", 10000, 0, "Intercompany expense"),
    ("090", "Business account", 25000, 0, "Cash at bank"),
    ("610", "Accounts receivable", 20000, 0, "Trade receivables"),
    ("710", "Office equipment", 15000, 0, "Fixed assets"),
    ("800", "Accounts payable", 0, 12000, "Trade payables"),
    ("820", "Owed to Company A", 0, 10000, "Intercompany payable"),
    ("970", "Owner funds", 0, 1000, "Share capital"),
    ("960", "Retained earnings", 0, 12000, "Retained earnings"),
]
ELIMS = [
    ("Remove management fee charged by A to B", "Intercompany income", 10000, 0),
    ("Remove management fee charged by A to B", "Intercompany expense", 0, 10000),
    ("Remove balance owed by B to A", "Intercompany payable", 10000, 0),
    ("Remove balance owed by B to A", "Intercompany receivable", 0, 10000),
]
COMPANIES = ["Company A", "Company B", "Company C"]


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HEAD, FILL
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w


def tb_sheet(wb, name, data, dv_ref):
    ws = wb.create_sheet(name)
    ws["A1"], ws["A1"].font = f"{name} - trial balance (replace with your export)", TITLE
    ws["A2"] = "Paste the trial balance for the period end. Yellow cells are inputs. Pick a group line for every account."
    header(ws, 4, ["Account code", "Account name", "Debit", "Credit", "Net (Dr - Cr)", "Group line"], [14, 36, 14, 14, 16, 28])
    dv = DataValidation(type="list", formula1=dv_ref, allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(5, 5 + ROWS):
        rec = data[r - 5] if r - 5 < len(data) else None
        if rec:
            for c, v in enumerate(rec[:4], 1):
                ws.cell(row=r, column=c, value=v)
            ws.cell(row=r, column=6, value=rec[4])
        ws.cell(row=r, column=5, value=f'=IF(AND(C{r}="",D{r}=""),"",N(C{r})-N(D{r}))').number_format = NUM
        for c in (1, 2, 3, 4, 6):
            ws.cell(row=r, column=c).fill = INPUT
        ws.cell(row=r, column=3).number_format = ws.cell(row=r, column=4).number_format = NUM
        dv.add(f"F{r}")
    end = 4 + ROWS
    t = end + 1
    ws.cell(row=t, column=2, value="Totals").font = Font(bold=True)
    for col in "CDE":
        ws[f"{col}{t}"] = f"=SUM({col}5:{col}{end})"
        ws[f"{col}{t}"].number_format, ws[f"{col}{t}"].font = NUM, Font(bold=True)
    ws.cell(row=t + 1, column=2, value="Check: balances?")
    ws[f"C{t + 1}"] = f'=IF(ROUND(C{t}-D{t},2)=0,"OK","Does not balance")'
    ws.cell(row=t + 2, column=2, value="Accounts with no group line")
    ws[f"C{t + 2}"] = f'=SUMPRODUCT((E5:E{end}<>"")*(F5:F{end}=""))'
    ws.freeze_panes = "A5"


def main():
    wb = Workbook()
    st = wb.active
    st.title = "Start Here"
    lines = [
        ("Multi-Entity Consolidation Template", TITLE),
        ("Free template from RankerToolAI (rankertoolai.com). Combine up to 3 companies' trial balances into a consolidated P&L and balance sheet.", None),
        ("", None),
        ("How to use it", Font(bold=True)),
        ("1. Group Lines: keep or edit the list of reporting lines (each has a statement and a section).", None),
        ("2. Company A / B / C: paste each company's trial balance for the same period end, in the same currency.", None),
        ("3. In column F of each company sheet, pick the group line every account belongs to (dropdown).", None),
        ("4. Eliminations: enter one line per side of every intercompany balance or transaction, so they net to zero.", None),
        ("5. Consolidated TB, P&L and Balance Sheet update automatically. Check that every 'Check' cell says OK.", None),
        ("", None),
        ("The example figures in Company A and Company B are illustrative: A charges B a 10,000 management fee and B still owes it.", None),
        ("Delete them and paste your own data. Company C is empty.", None),
        ("", None),
        ("Limits of a spreadsheet", Font(bold=True)),
        ("No foreign currency translation, no automatic refresh from Xero / QuickBooks / Sage / MYOB, and eliminations are manual each month.", None),
        ("When that becomes a monthly chore, see our comparison of consolidation software: https://rankertoolai.com/best/consolidation-software/", None),
        ("Cost by number of companies: https://rankertoolai.com/tools/consolidation-software-cost-calculator/", None),
    ]
    for i, (text, font) in enumerate(lines, 1):
        st.cell(row=i, column=1, value=text)
        if font:
            st.cell(row=i, column=1).font = font
    st.column_dimensions["A"].width = 130

    gl = wb.create_sheet("Group Lines")
    gl["A1"], gl["A1"].font = "Group reporting lines", TITLE
    header(gl, 3, ["Group line", "Statement", "Section"], [30, 12, 26])
    for i, (n, stmt, sec) in enumerate(GROUP_LINES, 4):
        for c, v in enumerate((n, stmt, sec), 1):
            gl.cell(row=i, column=c, value=v).fill = INPUT
    gl_end = 3 + len(GROUP_LINES) + 20  # room for extra lines
    dv_ref = f"'Group Lines'!$A$4:$A${gl_end}"

    tb_sheet(wb, "Company A", COMPANY_A, dv_ref)
    tb_sheet(wb, "Company B", COMPANY_B, dv_ref)
    tb_sheet(wb, "Company C", [], dv_ref)

    el = wb.create_sheet("Eliminations")
    el["A1"], el["A1"].font = "Intercompany eliminations", TITLE
    el["A2"] = "Each elimination needs two lines that cancel out (total Net must be 0)."
    header(el, 4, ["Description", "Group line", "Debit", "Credit", "Net (Dr - Cr)"], [44, 28, 14, 14, 16])
    dv = DataValidation(type="list", formula1=dv_ref, allow_blank=True)
    el.add_data_validation(dv)
    for r in range(5, 35):
        rec = ELIMS[r - 5] if r - 5 < len(ELIMS) else None
        if rec:
            for c, v in enumerate(rec, 1):
                el.cell(row=r, column=c, value=v)
        el.cell(row=r, column=5, value=f'=IF(AND(C{r}="",D{r}=""),"",N(C{r})-N(D{r}))').number_format = NUM
        for c in (1, 2, 3, 4):
            el.cell(row=r, column=c).fill = INPUT
        el.cell(row=r, column=3).number_format = el.cell(row=r, column=4).number_format = NUM
        dv.add(f"B{r}")
    el["B36"], el["E36"] = "Total net (should be 0)", "=SUM(E5:E34)"
    el["E36"].number_format = NUM
    el["B37"], el["E37"] = "Check", '=IF(ROUND(E36,2)=0,"OK","Eliminations do not net to zero")'

    ct = wb.create_sheet("Consolidated TB")
    ct["A1"], ct["A1"].font = "Consolidated trial balance (Net = Dr - Cr)", TITLE
    header(ct, 3, ["Group line", "Statement", "Section"] + COMPANIES + ["Eliminations", "Consolidated"], [30, 12, 24, 14, 14, 14, 14, 16])
    end = 4 + ROWS
    first, last = 4, 3 + len(GROUP_LINES) + 20
    for r in range(first, last + 1):
        i = r - first + 4
        ct[f"A{r}"] = f"=IF('Group Lines'!A{i}=\"\",\"\",'Group Lines'!A{i})"
        ct[f"B{r}"] = f"=IF('Group Lines'!B{i}=\"\",\"\",'Group Lines'!B{i})"
        ct[f"C{r}"] = f"=IF('Group Lines'!C{i}=\"\",\"\",'Group Lines'!C{i})"
        for j, comp in enumerate(COMPANIES):
            col = "DEF"[j]
            ct[f"{col}{r}"] = f"=IF(A{r}=\"\",\"\",SUMIF('{comp}'!$F$5:$F${end},A{r},'{comp}'!$E$5:$E${end}))"
        ct[f"G{r}"] = f"=IF(A{r}=\"\",\"\",SUMIF(Eliminations!$B$5:$B$34,A{r},Eliminations!$E$5:$E$34))"
        ct[f"H{r}"] = f"=IF(A{r}=\"\",\"\",SUM(D{r}:G{r}))"
        for col in "DEFGH":
            ct[f"{col}{r}"].number_format = NUM
    tot = last + 1
    ct[f"A{tot}"], ct[f"A{tot}"].font = "Total (should be 0)", Font(bold=True)
    for col in "DEFGH":
        ct[f"{col}{tot}"] = f"=SUM({col}{first}:{col}{last})"
        ct[f"{col}{tot}"].number_format, ct[f"{col}{tot}"].font = NUM, Font(bold=True)
    ct[f"A{tot + 1}"], ct[f"H{tot + 1}"] = "Check", f'=IF(ROUND(H{tot},2)=0,"OK","Does not balance")'
    ct.freeze_panes = "D4"
    rng = lambda c: f"'Consolidated TB'!${c}${first}:${c}${last}"

    def stmt_sheet(name, title, rows):
        ws = wb.create_sheet(name)
        ws["A1"], ws["A1"].font = title, TITLE
        header(ws, 3, [""] + COMPANIES + ["Eliminations", "Consolidated"], [34, 14, 14, 14, 14, 16])
        for i, (label, kind, arg, bold) in enumerate(rows, 4):
            ws[f"A{i}"] = label
            if bold:
                ws[f"A{i}"].font = Font(bold=True)
            for col, src in zip("BCDEF", "DEFGH"):
                if kind == "section":  # sign, section name
                    sign, sec = arg
                    ws[f"{col}{i}"] = f'={sign}SUMIF({rng("C")},"{sec}",{rng(src)})'
                elif kind == "formula":
                    ws[f"{col}{i}"] = "=" + arg.replace("#", col)
                if kind != "blank":
                    ws[f"{col}{i}"].number_format = NUM
                    if bold:
                        ws[f"{col}{i}"].font = Font(bold=True)
        return ws

    stmt_sheet("P&L", "Consolidated profit and loss", [
        ("Income", "section", ("-", "Income"), False),
        ("Cost of sales", "section", ("", "Cost of sales"), False),
        ("Gross profit", "formula", "#4-#5", True),
        ("Overheads", "section", ("", "Overheads"), False),
        ("Net profit", "formula", "#6-#7", True),
    ])
    stmt_sheet("Balance Sheet", "Consolidated balance sheet", [
        ("Non-current assets", "section", ("", "Non-current assets"), False),
        ("Current assets", "section", ("", "Current assets"), False),
        ("Total assets", "formula", "#4+#5", True),
        ("Current liabilities", "section", ("-", "Current liabilities"), False),
        ("Non-current liabilities", "section", ("-", "Non-current liabilities"), False),
        ("Net assets", "formula", "#6-#7-#8", True),
        ("Equity", "section", ("-", "Equity"), False),
        ("Profit for the period", "formula", "'P&L'!#8", False),
        ("Total equity", "formula", "#10+#11", True),
        ("Check (net assets = equity)", "formula", 'IF(ROUND(#9-#12,2)=0,"OK","Check")', False),
    ])
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title != "Start Here"
        for row in ws.iter_rows():
            for c in row:
                if c.alignment.horizontal is None and isinstance(c.value, str) and len(c.value) > 60:
                    c.alignment = Alignment(wrap_text=False)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
